#!/usr/bin/env python3
"""Build the encrypted CRM file (crm.enc) for the SUWEY sales app.

Usage:
  python3 tools/build_crm.py SALES_JSON STOCK_JSON PIN [COLLECTIONS_JSON]
  python3 tools/build_crm.py --from-excel XLSX OUT_SALES_JSON [OUT_COLLECTIONS_JSON]

  COLLECTIONS_JSON  the admin artifact's db document collections/current: TAHSİLAT RAPORU
                    (per customer: satis, nakit, havale, cek, kk, toplam, son, kalan, durum)
                    and TAHSİLAT KAYITLARI (payments: d, n, tutar, tur, not)

  SALES_JSON  the admin artifact's db document sales/current (bare body or {"data": ...});
              body: {"cols": [...], "rows": [[...], ...], "source", "updatedAt"}
  STOCK_JSON  the admin artifact's db document stock/current (for the product list)
  PIN         the app password (kept only in the admin artifact's db config/pin)

Discounts (İSKONTO) appear only inside proforma lines, never in the price list.
"""
import base64, collections, datetime, json, os, secrets, sys

from cryptography.hazmat.primitives import hashes
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ITER = 600_000
COLS = ["city", "cust", "p", "c", "s", "adet", "liste", "isk", "net", "tutar",
        "teslim", "satisT", "sevkT", "fatura", "odeme", "diger"]


def body(path):
    d = json.load(open(path, encoding="utf-8"))
    return d["data"] if isinstance(d, dict) and "data" in d and "rows" not in d else d


def num(v):
    try:
        return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0


def _h(v):
    return " ".join(str(v or "").split()).upper()


def _ymd(v):
    return v.strftime("%Y-%m-%d") if isinstance(v, (datetime.date, datetime.datetime)) else (str(v) if v else None)


def _sheet(wb, prefix):
    return next((wb[n] for n in wb.sheetnames if _h(n).startswith(prefix)), None)


def from_excel(xlsx, out_sales, out_coll=None):
    """Extract SATIŞ SUWEY rows (and TAHSİLAT RAPORU / KAYITLARI when out_coll is given)."""
    import openpyxl, warnings
    warnings.filterwarnings("ignore")
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    ws = _sheet(wb, "SATIŞ")
    rows, hdr = [], None
    for r in ws.iter_rows(values_only=True):
        if hdr is None:
            if _h(r[0]) == "ŞEHİR" and _h(r[1]) == "MÜŞTERİ":
                H = [_h(x) for x in r]
                f = lambda *names, d=None: next((i for i, h in enumerate(H) for n in names if h.startswith(n)), d)
                hdr = dict(city=0, cust=1, p=f("ÜRÜN ADI", d=2), c=f("RENK", d=3), s=f("BEDEN", d=4), adet=f("ADET", d=5),
                           liste=f("FİYAT", d=6), isk=f("İSKONTO", d=7), net=f("SATIŞ FİYATI", d=8), tutar=f("TOPLAM TUTAR", d=9),
                           teslim=f("TESLİM", d=10), satisT=f("SATIŞ TARİHİ", d=11), sevkT=f("SEVK TARİHİ", d=12),
                           fatura=f("FATURA TUTARI", d=15), odeme=f("TAHSİLAT DURUMU", "ÖDEME DURUMU", d=17),
                           diger=f("DİĞER TUTAR"))
            continue
        g = lambda k: r[hdr[k]] if hdr[k] is not None and hdr[k] < len(r) else None
        if not g("cust") or not g("p"):
            continue
        rows.append([str(g("city") or "").strip(), str(g("cust")).strip(), str(g("p")).strip(), str(g("c") or "").strip(),
                     str(g("s") or "").strip(), num(g("adet")), num(g("liste")), num(g("isk")), num(g("net")), num(g("tutar")),
                     str(g("teslim") or "").strip(), _ymd(g("satisT")), _ymd(g("sevkT")),
                     (None if g("fatura") in (None, "") else num(g("fatura"))), (str(g("odeme")).strip() if g("odeme") else None),
                     num(g("diger"))])
    now = datetime.datetime.now(datetime.timezone.utc).isoformat()
    json.dump({"cols": COLS, "rows": rows, "source": os.path.basename(xlsx), "updatedAt": now},
              open(out_sales, "w", encoding="utf-8"), ensure_ascii=False)
    print("sales rows", len(rows))
    if not out_coll:
        return
    coll = {"customers": [], "payments": [], "source": os.path.basename(xlsx), "updatedAt": now}
    ws = _sheet(wb, "TAHSİLAT RAPOR")
    if ws is not None:
        H = None
        for r in ws.iter_rows(values_only=True):
            if H is None:
                if _h(r[0]) == "MÜŞTERİ" and any(_h(x).startswith("KALAN") for x in r):
                    H = [_h(x) for x in r]
                    def ix(*names):
                        for n in names:
                            for i, h in enumerate(H):
                                if h.startswith(n):
                                    return i
                        raise KeyError("TAHSİLAT RAPORU column not found: " + " / ".join(names))
                    # satis = total sales incl. VAT; newer layout splits it into FATURA + DİĞER = TOPLAM SATIŞ
                    I = dict(satis=ix("TOPLAM SATIŞ", "SATIŞ TUTARI"), nakit=ix("NAKİT"), havale=ix("HAVALE"), cek=ix("ÇEK"), kk=ix("KREDİ"),
                             toplam=ix("TOPLAM TAHSİLAT"), son=ix("SON TAHSİLAT"), kalan=ix("KALAN"), durum=ix("TAHSİLAT DURUMU"))
                continue
            n = str(r[0] or "").strip()
            if not n or _h(n) == "TOPLAM" or n.startswith("Not"):
                if _h(n) == "TOPLAM":
                    break
                continue
            coll["customers"].append({"n": n, "satis": num(r[I["satis"]]), "nakit": num(r[I["nakit"]]), "havale": num(r[I["havale"]]),
                                      "cek": num(r[I["cek"]]), "kk": num(r[I["kk"]]), "toplam": num(r[I["toplam"]]),
                                      "son": _ymd(r[I["son"]]), "kalan": round(num(r[I["kalan"]]), 2),
                                      "durum": (str(r[I["durum"]]).strip() if r[I["durum"]] else None)})
    ws = _sheet(wb, "TAHSİLAT KAYIT")
    if ws is not None:
        started = False
        for r in ws.iter_rows(values_only=True):
            if not started:
                started = _h(r[0]) == "TARİH" and _h(r[1]) == "MÜŞTERİ"
                continue
            if not r[1] or r[2] in (None, ""):
                continue
            coll["payments"].append({"d": _ymd(r[0]), "n": str(r[1]).strip(), "tutar": num(r[2]),
                                     "tur": str(r[3] or "").strip(), "not": str(r[4] or "").strip()})
    json.dump(coll, open(out_coll, "w", encoding="utf-8"), ensure_ascii=False)
    print("collection customers", len(coll["customers"]), "payments", len(coll["payments"]))


def build(sales, stock, coll=None):
    ix = {c: i for i, c in enumerate(sales["cols"])}
    R = [dict((c, r[ix[c]] if c in ix else None) for c in COLS) for r in sales["rows"]]  # older data has no "diger"
    # price list: most common non-zero list price per product; ties -> most recent
    by_p = collections.defaultdict(list)
    for r in R:
        if r["liste"] > 0:
            by_p[r["p"]].append((r["satisT"] or "", r["liste"], r["adet"]))
    types = {row["p"]: row["t"] for row in stock["rows"]}
    prices = []
    for p in sorted(set(types) | set(by_p)):
        cnt = collections.Counter()
        last = {}
        for d, pr, a in by_p.get(p, []):
            cnt[pr] += a
            last[pr] = max(last.get(pr, ""), d)
        price = max(cnt, key=lambda k: (cnt[k], last[k])) if cnt else None
        prices.append({"p": p, "t": types.get(p, ""), "price": price})
    # customers -> orders (one order = customer + sale date)
    custs = {}
    for r in R:
        cu = custs.setdefault(r["cust"], {"n": r["cust"], "city": r["city"], "orders": {}})
        o = cu["orders"].setdefault(r["satisT"] or "", {"d": r["satisT"], "lines": []})
        o["lines"].append([r["p"], r["c"], r["s"], r["adet"], r["liste"], r["isk"], r["net"], r["tutar"],
                           r["teslim"], r["sevkT"], r["fatura"], r["odeme"], r["diger"] or 0])
    out = []
    for cu in custs.values():
        cu["orders"] = sorted(cu["orders"].values(), key=lambda o: o["d"] or "", reverse=True)
        out.append(cu)
    out.sort(key=lambda c: c["n"])
    return {"v": 1, "updatedAt": datetime.datetime.now(datetime.timezone.utc).isoformat(),
            "reportDate": stock.get("reportDate"), "prices": prices, "customers": out,
            "lineCols": ["p", "c", "s", "adet", "liste", "isk", "net", "tutar", "teslim", "sevkT", "fatura", "odeme", "diger"],
            "collections": ({"customers": coll.get("customers", []), "payments": coll.get("payments", [])} if coll else None)}


def encrypt(obj, pin):
    salt, iv = secrets.token_bytes(16), secrets.token_bytes(12)
    key = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=salt, iterations=ITER).derive(pin.encode())
    ct = AESGCM(key).encrypt(iv, json.dumps(obj, ensure_ascii=False, separators=(",", ":")).encode(), None)
    b = lambda x: base64.b64encode(x).decode()
    return {"v": 1, "kdf": "PBKDF2-SHA256", "iter": ITER, "salt": b(salt), "iv": b(iv), "ct": b(ct)}


if __name__ == "__main__":
    if sys.argv[1] == "--from-excel":
        from_excel(sys.argv[2], sys.argv[3], sys.argv[4] if len(sys.argv) > 4 else None)
        sys.exit()
    sales, stock, pin = body(sys.argv[1]), body(sys.argv[2]), sys.argv[3]
    coll = body(sys.argv[4]) if len(sys.argv) > 4 else None
    assert sales.get("rows"), "sales has no rows"
    crm = build(sales, stock, coll)
    json.dump(encrypt(crm, pin), open(os.path.join(ROOT, "crm.enc"), "w"))
    n_orders = sum(len(c["orders"]) for c in crm["customers"])
    print(f"customers={len(crm['customers'])} orders={n_orders} lines={len(sales['rows'])} prices={sum(1 for p in crm['prices'] if p['price'])}/{len(crm['prices'])}")
